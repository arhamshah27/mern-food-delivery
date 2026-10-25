import os
import re
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill_hex)
    tcPr.append(shd)

def create_word_document():
    doc = Document()
    
    # Page setup - 1 inch margins
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    # Styles & Fonts
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Calibri'
    normal_style.font.size = Pt(11)
    normal_style.font.color.rgb = RGBColor(30, 41, 59)

    # Document Title
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_title = title.add_run("Baroda Gourmet Kitchen")
    run_title.font.name = 'Calibri'
    run_title.font.size = Pt(24)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(234, 88, 12) # Deep orange

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_sub = subtitle.add_run("Full-Stack MERN Food Ordering & Real-Time Fleet Delivery Platform\nProject Submission & Evaluation Documentation")
    run_sub.font.name = 'Calibri'
    run_sub.font.size = Pt(13)
    run_sub.font.italic = True
    run_sub.font.color.rgb = RGBColor(71, 85, 105)

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # Read markdown file
    with open("PROJECT_SUBMISSION_DOC.md", "r", encoding="utf-8") as f:
        md_text = f.read()

    lines = md_text.split("\n")
    in_code_block = False
    code_content = []
    in_table = False
    table_rows = []

    for line in lines:
        stripped = line.strip()

        # Handle Code Blocks
        if stripped.startswith("```"):
            if in_code_block:
                # End of code block
                in_code_block = False
                p = doc.add_paragraph()
                p.paragraph_format.space_before = Pt(4)
                p.paragraph_format.space_after = Pt(8)
                p.paragraph_format.left_indent = Inches(0.2)
                code_run = p.add_run("\n".join(code_content))
                code_run.font.name = 'Consolas'
                code_run.font.size = Pt(9.5)
                code_run.font.color.rgb = RGBColor(15, 23, 42)
                code_content = []
            else:
                in_code_block = True
                code_content = []
            continue

        if in_code_block:
            code_content.append(line)
            continue

        # Handle Tables
        if "|" in line and (line.startswith("|") or stripped.startswith("|")):
            parts = [p.strip() for p in line.strip().split("|")[1:-1]]
            if all(re.match(r"^:?-+:?$", p) for p in parts):
                continue # Header separator
            if not in_table:
                in_table = True
                table_rows = [parts]
            else:
                table_rows.append(parts)
            continue
        else:
            if in_table:
                # Render table
                in_table = False
                if table_rows:
                    num_cols = max(len(r) for r in table_rows)
                    tbl = doc.add_table(rows=len(table_rows), cols=num_cols)
                    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
                    tbl.autofit = True
                    for r_idx, row_data in enumerate(table_rows):
                        row = tbl.rows[r_idx]
                        is_header = (r_idx == 0)
                        for c_idx, cell_value in enumerate(row_data):
                            if c_idx < num_cols:
                                cell = row.cells[c_idx]
                                cell.text = cell_value
                                for p in cell.paragraphs:
                                    p.paragraph_format.space_before = Pt(3)
                                    p.paragraph_format.space_after = Pt(3)
                                    for r in p.runs:
                                        r.font.name = 'Calibri'
                                        r.font.size = Pt(9.5 if is_header else 9)
                                        if is_header:
                                            r.font.bold = True
                                            r.font.color.rgb = RGBColor(255, 255, 255)
                                if is_header:
                                    set_cell_background(cell, "0F172A") # slate-900
                                elif r_idx % 2 == 1:
                                    set_cell_background(cell, "F8FAFC") # slate-50
                    doc.add_paragraph().paragraph_format.space_after = Pt(6)
                table_rows = []

        if not stripped:
            continue

        # Heading 1
        if stripped.startswith("# "):
            continue # Already handled in title

        # Heading 2
        if stripped.startswith("## "):
            h = doc.add_heading(level=1)
            h.paragraph_format.space_before = Pt(14)
            h.paragraph_format.space_after = Pt(4)
            h_run = h.add_run(stripped[3:].strip())
            h_run.font.name = 'Calibri'
            h_run.font.size = Pt(14)
            h_run.font.bold = True
            h_run.font.color.rgb = RGBColor(15, 23, 42)
            continue

        # Heading 3
        if stripped.startswith("### "):
            h = doc.add_heading(level=2)
            h.paragraph_format.space_before = Pt(10)
            h.paragraph_format.space_after = Pt(3)
            h_run = h.add_run(stripped[4:].strip())
            h_run.font.name = 'Calibri'
            h_run.font.size = Pt(12)
            h_run.font.bold = True
            h_run.font.color.rgb = RGBColor(51, 65, 85)
            continue

        # Heading 4
        if stripped.startswith("#### "):
            h = doc.add_heading(level=3)
            h.paragraph_format.space_before = Pt(8)
            h.paragraph_format.space_after = Pt(2)
            h_run = h.add_run(stripped[5:].strip())
            h_run.font.name = 'Calibri'
            h_run.font.size = Pt(11)
            h_run.font.bold = True
            h_run.font.color.rgb = RGBColor(71, 85, 105)
            continue

        # Bullet list items
        if stripped.startswith("- ") or stripped.startswith("* "):
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(2)
            p.add_run(stripped[2:])
            continue

        # Numbered list items
        num_match = re.match(r"^(\d+)\.\s+(.*)$", stripped)
        if num_match:
            p = doc.add_paragraph(style='List Number')
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            p.add_run(num_match.group(2))
            continue

        # Normal paragraph
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(4)
        p.add_run(stripped)

    # If file ended inside table
    if in_table and table_rows:
        num_cols = max(len(r) for r in table_rows)
        tbl = doc.add_table(rows=len(table_rows), cols=num_cols)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        for r_idx, row_data in enumerate(table_rows):
            row = tbl.rows[r_idx]
            is_header = (r_idx == 0)
            for c_idx, cell_value in enumerate(row_data):
                if c_idx < num_cols:
                    cell = row.cells[c_idx]
                    cell.text = cell_value
                    if is_header:
                        set_cell_background(cell, "0F172A")
        doc.add_paragraph().paragraph_format.space_after = Pt(6)

    output_filename = "Baroda_Gourmet_Project_Submission_Document.docx"
    doc.save(output_filename)
    print(f"Generated {output_filename} successfully!")

if __name__ == "__main__":
    create_word_document()
